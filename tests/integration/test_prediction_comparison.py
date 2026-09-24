"""Offline comparison joins preserve source meaning and never rerun inference."""

import json
from unittest.mock import Mock

import pandas as pd
import pytest

from pedflow.exceptions import ContractError
from pedflow.prediction.comparison import (
    AUDIT_COLUMNS,
    BAND_COLUMNS,
    FEATURE_COLUMNS,
    OBS_COLUMNS,
    OUTPUT_COLUMNS,
    add_comparisons,
)
from pedflow.prediction.io import refresh_prediction_outputs, write_prediction_outputs
from pedflow.prediction.workflow import PredictionResult


@pytest.fixture
def comparison_case(tmp_path):
    """Use deliberately different labels, numeric references and corrected observations."""
    for folder in ("raw", "processed"):
        (tmp_path / "data" / folder).mkdir(parents=True)
    sites = pd.DataFrame(
        {
            "site_id": ["001", "NA"],
            "latitude": [-32.0, -31.0],
            "longitude": [116.0, 115.0],
            "coordinate_crs": ["EPSG:4326"] * 2,
        }
    )
    predictions = sites.assign(
        pred_peak_cross_ped_per_h=[1.5, 2.5],
        pred_peak_along_ped_per_h=[3.5, 4.5],
        status="predicted",
        message="",
    )
    model = sites.iloc[:1].assign(
        source_location_id="old-road-id",
        obs_peak_cross_ped_per_h=87,
        obs_peak_along_ped_per_h=17.3333,
    )
    model.to_excel(
        tmp_path / "data/processed/model_dataset.xlsx", sheet_name="Model_Data", index=False
    )
    links = sites.drop(columns="coordinate_crs").assign(
        road_no=["H001", "H002"],
        slk_km=1.5,
        road_name="Road",
        local_government_name="Town",
        matched_segment_id=["AUSRAP-ROW-000002", "AUSRAP-ROW-000003"],
        irap_source_excel_row=[2, 3],
    )
    link_path = tmp_path / "data/processed/comparison_sites.csv"
    links.to_csv(link_path, index=False)
    raw = pd.DataFrame(
        {
            "Road No": ["H001", "H002"],
            "Start SLK": [1.0, 1.0],
            "End SLK": [2.0, 2.0],
            "Road survey date": ["2023-06-01", "2023-06-02"],
            **{column: ["2 to 3 pedestrians", None] for column in BAND_COLUMNS},
        }
    )
    raw_path = tmp_path / "data/raw/AusRAP Data - Words - MSDS.xlsx"
    raw.to_excel(raw_path, index=False)
    pd.DataFrame(
        {
            "ID": ["old-road-id"],
            "Ped_Peak_Cross": [0.0],
            "Ped_Peak_Along": [18.0],
            "Ped_Peak_Driver": [15.0],
            "Ped_Peak_Passenger": [3.0],
            "obs_peak_cross": [125],
        }
    ).to_csv(tmp_path / "data/raw/Ped_iRAP_merged.csv", index=False)
    return tmp_path, sites, predictions


def test_raw_numeric_zero_labels_and_corrected_observations_are_distinct(comparison_case):
    root, sites, predictions = comparison_case
    output, audit, metadata = add_comparisons(predictions, sites, root)
    assert tuple(output.columns) == OUTPUT_COLUMNS
    assert output.site_id.tolist() == ["001", "NA"]
    assert output.obs_peak_cross_ped_per_h.iloc[0] == 87  # Never the client's old 125.
    assert output.irap_peak_cross_reference.iloc[0] == 0  # Preserve a real zero.
    assert output.irap_peak_along_reference.iloc[0] == 18
    assert output.irap_peak_driver_reference.iloc[0] == 15
    assert output.irap_peak_passenger_reference.iloc[0] == 3
    assert output.matched_segment_id.iloc[0] == "AUSRAP-ROW-000002"
    assert output.irap_numeric_source_id.iloc[0] == "old-road-id"
    assert metadata["site_records"][0]["irap_source_excel_row"] == 2
    assert output.irap_ped_cross_band.iloc[0] == "2 to 3 pedestrians"
    assert (
        output.loc[1, [*OBS_COLUMNS, "irap_peak_cross_reference", "irap_ped_cross_band"]]
        .isna()
        .all()
    )
    assert output.location_role.tolist() == ["observed_benchmark", "unobserved"]
    assert audit.irap_match_status.eq("matched").all()  # A blank source label is not no-match.
    assert audit.irap_numeric_status.tolist() == ["matched", "no_reference"]
    assert all(
        source["status"] == "loaded" and "sha256" not in source for source in metadata["sources"]
    )


@pytest.mark.parametrize("change", ["id", "coordinates", "road", "slk"])
def test_reused_ids_or_conflicting_road_hints_do_not_attach_wrong_segment(comparison_case, change):
    root, sites, predictions = comparison_case
    if change == "id":
        sites.loc[0, "site_id"] = "unknown"
        predictions.loc[0, "site_id"] = "unknown"
    elif change == "coordinates":
        sites.loc[0, "latitude"] = -30
    elif change == "road":
        sites["road_no"] = ["wrong", "H002"]
    else:
        sites["slk_km"] = [9, 1.5]
    output, audit, _ = add_comparisons(predictions, sites, root)
    assert pd.isna(output.irap_ped_cross_band.iloc[0])
    assert audit.irap_match_status.iloc[0] != "matched"
    assert output.pred_peak_cross_ped_per_h.tolist() == [1.5, 2.5]
    if change in {"id", "coordinates"}:
        assert pd.isna(output.obs_peak_cross_ped_per_h.iloc[0])


def test_updated_source_labels_are_loaded_without_a_fingerprint_manifest(comparison_case):
    root, sites, predictions = comparison_case
    path = root / "data/raw/AusRAP Data - Words - MSDS.xlsx"
    raw = pd.read_excel(path)
    raw.loc[0, next(iter(BAND_COLUMNS))] = "6 to 10 pedestrians"
    raw.to_excel(path, index=False)
    output, audit, _ = add_comparisons(predictions, sites, root)
    assert output.irap_ped_cross_band.iloc[0] == "6 to 10 pedestrians"
    assert audit.irap_match_status.eq("matched").all()
    assert output.pred_peak_cross_ped_per_h.tolist() == [1.5, 2.5]


def test_literal_none_label_is_not_treated_as_a_missing_cell(comparison_case):
    """The word export contains the actual category text 'None'; retain it verbatim."""
    root, sites, predictions = comparison_case
    path = root / "data/raw/AusRAP Data - Words - MSDS.xlsx"
    raw = pd.read_excel(path)
    raw.loc[0, next(iter(BAND_COLUMNS))] = "None"
    raw.to_excel(path, index=False)
    output, _, _ = add_comparisons(predictions, sites, root)
    assert output.irap_ped_cross_band.iloc[0] == "None"
    assert pd.isna(output.irap_ped_cross_band.iloc[1])


def test_comparison_join_preserves_shuffled_input_order(comparison_case):
    root, sites, predictions = comparison_case
    output, audit, _ = add_comparisons(predictions.iloc[::-1], sites.iloc[::-1], root)
    assert output.site_id.tolist() == ["NA", "001"]
    assert audit.irap_source_excel_row.tolist() == [3, 2]
    assert output.obs_peak_cross_ped_per_h.iloc[1] == 87


def test_refresh_clears_old_references_when_sources_disappear(comparison_case):
    """A previously populated report must not disguise stale values as current matches."""
    root, sites, predictions = comparison_case
    previous, _, _ = add_comparisons(predictions, sites, root)
    (root / "data/raw/Ped_iRAP_merged.csv").unlink()
    (root / "data/raw/AusRAP Data - Words - MSDS.xlsx").unlink()
    (root / "data/processed/model_dataset.xlsx").unlink()
    output, audit, _ = add_comparisons(previous, sites, root)
    assert (
        output[
            [
                *OBS_COLUMNS,
                *BAND_COLUMNS.values(),
                "irap_peak_cross_reference",
                "irap_peak_along_reference",
            ]
        ]
        .isna()
        .all()
        .all()
    )
    assert audit.benchmark_match_status.eq("source_unavailable").all()
    assert output.pred_peak_cross_ped_per_h.tolist() == [1.5, 2.5]


def test_offline_refresh_preserves_run_features_and_predictions(comparison_case, monkeypatch):
    root, sites, predictions = comparison_case
    features = sites[["site_id"]].assign(commercial_buildings=[1, 2], footpath_both_sides=True)
    metadata = {
        "run_id": "original-run",
        "configuration": {"model_version": "original-model"},
        "started_at_utc": "original-start",
        "finished_at_utc": "original-finish",
    }
    folder = root / "output"
    write_prediction_outputs(PredictionResult(predictions, features, metadata), folder)
    forbidden = Mock(side_effect=AssertionError("No inference or provider calls during refresh"))
    monkeypatch.setattr("pedflow.prediction.__main__.predict_sites", forbidden)
    from pedflow.prediction.__main__ import main

    monkeypatch.setattr(
        "sys.argv",
        [
            "pedflow-predict",
            "--project-root",
            str(root),
            "--output-directory",
            str(folder),
            "--refresh-comparison",
        ],
    )
    assert main() == 0
    forbidden.assert_not_called()
    saved = pd.read_excel(folder / "predictions.xlsx", sheet_name=None, keep_default_na=False)
    assert list(saved) == ["Predictions", "Features", "Metadata"]
    assert saved["Predictions"].pred_peak_cross_ped_per_h.tolist() == [1.5, 2.5]
    assert "coordinate_crs" not in saved["Predictions"]
    assert "model_version" not in saved["Predictions"]
    assert saved["Features"].commercial_buildings.tolist() == [1, 2]
    assert tuple(saved["Features"].columns) == FEATURE_COLUMNS
    assert not set(AUDIT_COLUMNS).intersection(saved["Features"].columns)
    assert saved["Predictions"].irap_peak_driver_reference.iloc[0] == 15
    from openpyxl import load_workbook

    styled = load_workbook(folder / "predictions.xlsx")["Predictions"]
    assert styled["A1"].fill.patternType is None
    assert styled["A1"].font.color.rgb == "00000000"
    assert styled["A2"].border.bottom.color.rgb == "00D9D9D9"
    manifest = json.loads((folder / "predictions.manifest.json").read_text())
    assert manifest["run_id"] == "original-run"
    assert manifest["finished_at_utc"] == "original-finish"
    assert manifest["model_version"] == "original-model"
    assert "workbook_sha256" not in manifest
    assert not saved["Metadata"].field.str.contains("sha256").any()
    # Refresh is repeatable and must not duplicate comparison/audit columns.
    refresh_prediction_outputs(folder, root, sites=sites.iloc[::-1])
    again = pd.read_excel(
        folder / "predictions.xlsx", sheet_name="Predictions", keep_default_na=False
    )
    pd.testing.assert_frame_equal(saved["Predictions"], again)
    bad = sites.copy()
    bad.loc[0, "latitude"] = -30
    with pytest.raises(ContractError, match="coordinates"):
        refresh_prediction_outputs(folder, root, sites=bad)
    # A formatting edit must not prevent an offline comparison refresh.
    workbook = load_workbook(folder / "predictions.xlsx")
    workbook["Predictions"].column_dimensions["A"].width = 30
    workbook.save(folder / "predictions.xlsx")
    refresh_prediction_outputs(folder, root)
    edited = pd.read_excel(folder / "predictions.xlsx", keep_default_na=False)
    pd.testing.assert_frame_equal(saved["Predictions"], edited)
