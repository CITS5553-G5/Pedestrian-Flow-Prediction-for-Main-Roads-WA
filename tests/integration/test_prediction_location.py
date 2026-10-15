"""Location matching and v5 output preserve model and reference distinctions."""

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from pedflow.exceptions import ContractError
from pedflow.io import read_sites
from pedflow.prediction.io import write_prediction_outputs
from pedflow.prediction.location import resolve_locations
from pedflow.prediction.schema import FEATURE_COLUMNS, OUTPUT_COLUMNS
from pedflow.prediction.workflow import PredictionResult

ROOT = Path(__file__).resolve().parents[2]


@pytest.mark.parametrize("path", [
    "data/processed/batch_input_dataset.xlsx", "data/templates/batch_input_template.xlsx",
])
def test_example_batches_preserve_observations_and_location_ambiguities(path):
    source = read_sites(ROOT / path)
    assert len(source) > 0
    assert source.site_id.is_unique
    assert set(source.site_group) == {"random", "observed"}
    observed = source.loc[source.site_group.eq("observed")].reset_index(drop=True)
    model = pd.read_excel(ROOT / "data/processed/model_dataset.xlsx", sheet_name="Model_Data")
    fields = ["site_id", "latitude", "longitude", "road_no", "slk_km",
              "obs_peak_cross_ped_per_h", "obs_peak_along_ped_per_h"]
    expected = model.set_index("site_id").loc[observed.site_id].reset_index()
    pd.testing.assert_frame_equal(observed[fields], expected[fields], check_dtype=False)
    assert source.loc[source.site_group.eq("random"), fields[-2:]].eq("").all().all()
    matched = resolve_locations(source)
    assert len(matched) == len(source)
    assert matched.location_status.tolist() == source.location_precheck.tolist()
    assert matched.loc[source.site_group.eq("random"), "location_status"].eq("matched").all()
    assert matched.loc[matched.location_status.eq("matched"), "road_offset_m"].max() < 30


def test_new_road_slk_and_coordinate_inputs_match_both_directions():
    point = resolve_locations(pd.DataFrame({"site_id": ["unregistered-wa-road"], "road_no": ["H001"], "slk_km": [356.36]}))
    assert point.location_status.iloc[0] == "matched"
    assert point.matched_road_no.iloc[0] == "H001"
    reverse = resolve_locations(point[["site_id", "latitude", "longitude"]])
    assert reverse.location_status.iloc[0] == "matched"
    assert reverse.matched_road_no.iloc[0] == "H001"
    assert reverse.matched_slk_km.iloc[0] == pytest.approx(356.36, abs=0.01)


def test_official_segment_uses_slk_midpoint_on_road():
    segment = resolve_locations(pd.DataFrame({
        "site_id": ["official-row"], "road_no": ["H001"],
        "start_slk": [356.30], "end_slk": [356.40],
    }))
    assert segment.input_type.iloc[0] == "road_segment"
    assert segment.location_status.iloc[0] == "matched"
    assert segment.slk_km.iloc[0] == pytest.approx(356.35)
    point = resolve_locations(pd.DataFrame({"site_id": ["point"], "road_no": ["H001"], "slk_km": [356.35]}))
    assert segment.latitude.iloc[0] == pytest.approx(point.latitude.iloc[0])
    assert segment.longitude.iloc[0] == pytest.approx(point.longitude.iloc[0])


def test_divided_road_requires_carriageway_when_ambiguous():
    locations = resolve_locations(pd.DataFrame({
        "site_id": ["unspecified", "left"], "road_no": ["8140043", "8140043"],
        "slk_km": [0.2, 0.2], "cway": [None, "Left"],
    }))
    assert locations.location_status.tolist() == ["ambiguous_carriageway", "matched"]
    assert pd.isna(locations.latitude.iloc[0])
    assert locations.matched_cway.iloc[1] == "Left"


def test_both_inputs_use_entered_coordinates_and_flag_conflict():
    first = resolve_locations(pd.DataFrame({"site_id": ["a"], "road_no": ["H001"], "slk_km": [356.36]}))
    other = resolve_locations(pd.DataFrame({"site_id": ["b"], "road_no": ["H001"], "slk_km": [357.36]}))
    dual = pd.DataFrame({
        "site_id": ["dual"], "road_no": ["H001"], "slk_km": [356.36],
        "latitude": [other.latitude.iloc[0]], "longitude": [other.longitude.iloc[0]],
    })
    resolved = resolve_locations(dual)
    assert resolved.location_status.iloc[0] == "road_coordinate_conflict"
    assert resolved.latitude.iloc[0] == other.latitude.iloc[0]
    assert resolved.latitude.iloc[0] != first.latitude.iloc[0]


def test_invalid_incomplete_location_fails_clearly():
    with pytest.raises(ContractError, match="together"):
        resolve_locations(pd.DataFrame({"site_id": ["x"], "latitude": [-32.0]}))
    with pytest.raises(ContractError, match="road_no"):
        resolve_locations(pd.DataFrame({"site_id": ["x"], "road_no": ["H001"]}))


def test_export_has_two_predictions_and_nullable_reference_columns(tmp_path):
    prediction = pd.DataFrame([{
        "site_id": "new-location", "input_type": "coordinates", "latitude": -32.0,
        "longitude": 116.0, "location_status": "no_nearby_road",
        "pred_peak_cross_ped_per_h": 12.0, "pred_peak_along_ped_per_h": 15.0,
        "status": "predicted_with_warning", "message": "road:no_nearby_road",
    }]).reindex(columns=OUTPUT_COLUMNS)
    features = pd.DataFrame([{"site_id": "new-location", "commercial_buildings": 1,
                              "footpath_both_sides": True, "abs_mb_category_simple": "Other"}]).reindex(columns=FEATURE_COLUMNS)
    result = PredictionResult(prediction, features, {"run_id": "test", "status_counts": {"predicted_with_warning": 1}})
    paths = write_prediction_outputs(result, tmp_path)
    book = load_workbook(paths["workbook"])
    assert book.sheetnames == ["Metadata", "Prediction_Output", "Features"]
    assert tuple(c.value for c in book["Prediction_Output"][1][:len(OUTPUT_COLUMNS)]) == OUTPUT_COLUMNS
    assert len(OUTPUT_COLUMNS) == 29
    assert book["Prediction_Output"]["O2"].value is None
    assert book["Prediction_Output"]["P2"].value == 12.0
    assert book["Prediction_Output"]["Q2"].value is None
    assert book["Prediction_Output"]["W2"].value == 15.0
    assert book["Prediction_Output"]["X2"].value is None
    assert not any("mae" in c.lower() or "rmse" in c.lower() for c in OUTPUT_COLUMNS)
    assert book["Metadata"]["B14"].value == "prediction_output_v5"
    assert book["Metadata"]["A47"].value == "comparison_notes"
