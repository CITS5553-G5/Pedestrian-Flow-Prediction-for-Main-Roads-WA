"""Exercise real orchestration/model/checkpoints with offline source adapters."""

import json
from dataclasses import replace
from pathlib import Path
from unittest.mock import Mock

import pandas as pd
import pytest
from openpyxl import load_workbook

from pedflow.context.contracts import FeatureResult
from pedflow.context.sources import abs_meshblock_sa1
from pedflow.exceptions import ConfigurationError, ContractError
from pedflow.imagery.contracts import ExtractionError
from pedflow.modeling import load_model
from pedflow.prediction import predict_sites
from pedflow.prediction.io import write_prediction_outputs

ROOT = Path(__file__).resolve().parents[2]


@pytest.fixture
def sites():
    return pd.DataFrame(
        {
            "site_id": ["003", "001", "002"],
            "latitude": [-32.0] * 3,
            "longitude": [116.0] * 3,
            "coordinate_crs": ["EPSG:4326"] * 3,
        },
        index=[9, 1, 7],
    )


@pytest.fixture
def providers(monkeypatch):
    table = pd.DataFrame(
        {
            "site_id": ["001", "002", "003"],
            "abs_mb_category_2021": ["Commercial", "Parkland", "Residential"],
            "abs_match_status": ["matched"] * 3,
        }
    )
    source = Mock(
        side_effect=lambda *args: FeatureResult(
            "abs_meshblock_sa1", "offline fixture", table.copy(), {"fixture": True}
        )
    )
    monkeypatch.setattr(
        abs_meshblock_sa1, "DEFINITION", replace(abs_meshblock_sa1.DEFINITION, extractor=source)
    )
    extractor = Mock(
        side_effect=lambda site: (
            {"commercial_buildings": int(site["site_id"]), "footpath_both_sides": True},
            {"image_count": 6, "warnings": [], "source": "offline fixture"},
        )
    )
    factory = Mock(return_value=extractor)
    monkeypatch.setattr("pedflow.imagery.workflow.GoogleGeminiExtractor", factory)
    return table, source, extractor, factory


def test_order_mapping_and_prediction_parity(sites, providers):
    result = predict_sites(sites, ROOT, workers=2)
    assert result.predictions.site_id.tolist() == ["003", "001", "002"]
    assert result.features.abs_mb_category_simple.tolist() == ["Residential", "Commercial", "Other"]
    expected = load_model().predict(result.features)
    for name in ("pred_peak_cross_ped_per_h", "pred_peak_along_ped_per_h"):
        pd.testing.assert_series_equal(result.predictions[name], expected[name])
    assert result.predictions.status.eq("predicted_with_warning").all()
    assert result.metadata["context"]["sources"][0]["metadata"] == {"fixture": True}
    assert result.metadata["imagery"]["configuration"]["max_workers"] == 2


def test_road_only_input_predicts_from_measured_coordinate(providers):
    providers[0].drop(index=[1, 2], inplace=True)
    source = pd.DataFrame({"site_id": ["001"], "road_no": ["H001"], "slk_km": [356.36]})
    result = predict_sites(source, ROOT)
    row = result.predictions.iloc[0]
    assert row.input_type == "road_slk"
    assert row.location_status == "matched"
    assert row.matched_road_no == "H001"
    assert row.latitude < 0
    assert row.longitude > 0
    assert row.status == "predicted"
    assert pd.notna(row.pred_peak_cross_ped_per_h)


def test_current_50_site_batch_reaches_model_and_export(providers, tmp_path):
    source = pd.read_excel(ROOT / "data/processed/batch_input_dataset.xlsx", sheet_name="Batch_Input").head(50)
    table = providers[0]
    table.drop(index=table.index, inplace=True)
    for site_id in source.site_id:
        table.loc[len(table)] = [site_id, "Commercial", "matched"]
    providers[2].side_effect = lambda site: (
        {"commercial_buildings": 1, "footpath_both_sides": True},
        {"image_count": 6, "warnings": [], "source": "offline fixture"},
    )
    result = predict_sites(source, ROOT)
    assert len(result.predictions) == 50
    assert result.predictions.location_status.eq("matched").all()
    assert result.predictions.status.eq("predicted").all()
    assert result.predictions[list(("pred_peak_cross_ped_per_h", "pred_peak_along_ped_per_h"))].notna().all().all()
    output = write_prediction_outputs(result, tmp_path)["workbook"]
    saved = pd.read_excel(output, sheet_name="Prediction_Output")
    assert saved.site_id.tolist() == source.site_id.tolist()


def test_unresolved_road_row_does_not_block_other_locations(providers):
    providers[0].drop(index=[1, 2], inplace=True)
    source = pd.DataFrame({
        "site_id": ["002", "001"], "road_no": ["UNKNOWN", None], "slk_km": [1.0, None],
        "latitude": [None, -32.0], "longitude": [None, 116.0],
    })
    result = predict_sites(source, ROOT)
    assert result.predictions.site_id.tolist() == ["002", "001"]
    assert result.predictions.status.tolist() == ["not_predicted", "predicted_with_warning"]
    assert pd.isna(result.predictions.pred_peak_cross_ped_per_h.iloc[0])
    assert result.features.site_id.tolist() == ["002", "001"]


def test_shuffled_imagery_results_join_by_id(sites, providers, monkeypatch):
    from pedflow.imagery.workflow import build_imagery_features

    expected = predict_sites(sites, ROOT)

    def shuffled(*args, **kwargs):
        result = build_imagery_features(*args, **kwargs)
        result.features = result.features.iloc[::-1]
        return result

    monkeypatch.setattr("pedflow.prediction.workflow.build_imagery_features", shuffled)
    actual = predict_sites(sites, ROOT)
    pd.testing.assert_frame_equal(actual.predictions, expected.predictions)


def test_partial_failures_remain_null_and_resume_retries_only_failed(sites, providers, tmp_path):
    table, source, extractor, _ = providers
    table.loc[0, "abs_match_status"] = "request_failed"
    original = extractor.side_effect

    def first(site):
        if site["site_id"] == "002":
            raise ExtractionError("llm_http_503")
        return original(site)

    extractor.side_effect = first
    checkpoint = tmp_path / "imagery.json"
    result = predict_sites(sites, ROOT, checkpoint_path=checkpoint)
    assert result.predictions.status.tolist() == ["predicted_with_warning", "not_predicted", "not_predicted"]
    assert result.predictions.pred_peak_cross_ped_per_h.iloc[1:].isna().all()
    assert result.predictions.message.str.contains("request_failed").any()
    table.loc[0, "abs_match_status"] = "matched"
    extractor.reset_mock()
    extractor.side_effect = original
    resumed = predict_sites(sites.iloc[::-1], ROOT, checkpoint_path=checkpoint, resume=True)
    extractor.assert_called_once()
    assert extractor.call_args.args[0]["site_id"] == "002"
    assert source.call_count == 2  # ABS is refreshed, including recovered matches.
    assert resumed.predictions.status.eq("predicted_with_warning").all()
    assert resumed.metadata["imagery"]["reused_sites"] == 2


def test_missing_category_is_not_other_and_unused_counts_do_not_block(sites, providers):
    table, _, _, _ = providers
    table.loc[0, "abs_mb_category_2021"] = None
    table.loc[2, "abs_match_status"] = "matched_no_local_counts"
    result = predict_sites(sites, ROOT)
    assert pd.isna(result.features.abs_mb_category_simple.iloc[1])
    assert result.predictions.status.tolist() == [
        "predicted_with_warning",
        "not_predicted",
        "predicted_with_warning",
    ]
    assert pd.isna(result.predictions.pred_peak_cross_ped_per_h.iloc[1])


def test_imagery_warning_missing_value_and_overflow_are_isolated(sites, providers):
    _, _, extractor, _ = providers
    extractor.side_effect = lambda site: (
        {
            "commercial_buildings": {"003": 13, "001": None, "002": 10**12}[site["site_id"]],
            "footpath_both_sides": True,
        },
        {"warnings": ["aerial_only"]},
    )
    result = predict_sites(sites, ROOT)
    assert result.predictions.status.tolist() == [
        "predicted_with_warning",
        "not_predicted",
        "prediction_failed",
    ]
    assert "aerial_only" in result.predictions.message.iloc[0]
    assert "outside training range" in result.predictions.message.iloc[0]
    assert result.predictions.pred_peak_along_ped_per_h.iloc[1:].isna().all()


@pytest.mark.parametrize("bad", ["empty", "duplicate", "coordinate"])
def test_invalid_input_fails_before_extraction(sites, providers, bad):
    if bad == "empty":
        sites = sites.iloc[:0]
    elif bad == "duplicate":
        sites.iloc[1, 0] = "003"
    else:
        sites.iloc[0, 1] = 999
    with pytest.raises(ContractError):
        predict_sites(sites, ROOT)
    providers[1].assert_not_called()
    providers[3].assert_not_called()


@pytest.mark.parametrize(
    "field", ["model_release_sha256", "imagery_recipe_sha256", "context_feature_set_version"]
)
def test_incompatible_configuration_fails_before_sources(sites, providers, tmp_path, field):
    settings = json.loads((ROOT / "configs/prediction/workflow_v1.json").read_text())
    settings[field] = "changed"
    config = tmp_path / "bad.json"
    config.write_text(json.dumps(settings))
    with pytest.raises(ConfigurationError):
        predict_sites(sites, ROOT, config_path=config)
    providers[1].assert_not_called()
    providers[3].assert_not_called()


@pytest.mark.parametrize("bad", ["missing", "extra", "duplicate"])
def test_broken_source_coverage_never_silently_drops_sites(sites, providers, bad):
    table, _, _, factory = providers
    if bad == "missing":
        table.drop(index=0, inplace=True)
    elif bad == "extra":
        table.loc[3] = ["unexpected", "Commercial", "matched"]
    else:
        table.loc[1, "site_id"] = "001"
    with pytest.raises(ContractError):
        predict_sites(sites, ROOT)
    factory.assert_not_called()


def test_cli_workbook_manifest_and_resume_match_service(sites, providers, tmp_path, monkeypatch):
    from pedflow.prediction.__main__ import main

    source = tmp_path / "sites.xlsx"
    sites.drop(columns="coordinate_crs").to_excel(source, sheet_name="Batch_Input", index=False)
    output = tmp_path / "output"
    argv = [
        "pedflow-predict",
        "--project-root",
        str(ROOT),
        "--input",
        str(source),
        "--output-directory",
        str(output),
        "--workers",
        "1",
    ]
    monkeypatch.setattr("sys.argv", argv)
    assert main() == 0
    workbook = output / "predictions.xlsx"
    saved = pd.read_excel(workbook, sheet_name="Prediction_Output", dtype={"site_id": "string"})
    expected = predict_sites(sites, ROOT).predictions
    assert saved.site_id.tolist() == expected.site_id.tolist()
    for name in ("pred_peak_cross_ped_per_h", "pred_peak_along_ped_per_h"):
        pd.testing.assert_series_equal(saved[name], expected[name])
    assert load_workbook(workbook).sheetnames == ["Metadata", "Prediction_Output", "Features"]
    manifest = json.loads((output / "predictions.manifest.json").read_text())
    assert "workbook_sha256" not in manifest
    assert manifest["status_counts"] == {"predicted_with_warning": 3}
    providers[3].reset_mock()
    monkeypatch.setattr("sys.argv", argv + ["--resume"])
    assert main() == 0
    providers[3].assert_not_called()


def test_cli_returns_partial_status(sites, providers, tmp_path, monkeypatch):
    from pedflow.prediction.__main__ import main

    providers[0].loc[0, "abs_match_status"] = "no_match"
    source = tmp_path / "sites.csv"
    sites.to_csv(source, index=False)
    monkeypatch.setattr(
        "sys.argv",
        [
            "pedflow-predict",
            "--project-root",
            str(ROOT),
            "--input",
            str(source),
            "--output-directory",
            str(tmp_path / "output"),
        ],
    )
    assert main() == 2


def test_excel_site_id_is_literal_text(sites, providers, tmp_path):
    result = predict_sites(sites, ROOT)
    result.predictions.loc[0, "site_id"] = '=HYPERLINK("example")'
    paths = write_prediction_outputs(result, tmp_path)
    cell = load_workbook(paths["workbook"])["Prediction_Output"]["A2"]
    assert cell.data_type == "s"
    assert cell.value == '=HYPERLINK("example")'


@pytest.fixture
def small_ausrap(tmp_path):
    from pedflow.prediction.ausrap import SOURCE_COLUMNS, AusRAPIndex

    source = tmp_path / "reference.xlsx"
    pd.DataFrame([
        ["H001", 356.3, 356.4, "S", "None", "1 pedestrian", "2 to 3 pedestrians"],
        ["H001", 356.4, 356.5, "S", "2 to 3 pedestrians", "None", "None"],
    ], columns=SOURCE_COLUMNS).to_excel(source, index=False)
    return AusRAPIndex(source, tmp_path / "cache")


def test_prediction_and_reference_share_one_export_without_accuracy_metrics(providers, small_ausrap, tmp_path):
    providers[0].drop(index=[1, 2], inplace=True)
    source = pd.DataFrame({"site_id": ["001"], "road_no": ["H001"], "slk_km": [356.36]})
    result = predict_sites(source, ROOT, ausrap_index=small_ausrap)
    row = result.predictions.iloc[0]
    expected = load_model().predict(result.features).iloc[0]
    assert row.pred_peak_cross_ped_per_h == expected.pred_peak_cross_ped_per_h
    assert row.pred_peak_along_ped_per_h == expected.pred_peak_along_ped_per_h
    assert row.existing_irap_cross_recoded_ped_per_h == 7.5
    assert row.existing_irap_along_recoded_ped_per_h == 9
    assert row.comparison_status == "matched"
    assert pd.isna(row.obs_peak_cross_ped_per_h)
    assert pd.isna(row.existing_ausrap_cross_category)
    assert result.metadata["ausrap"]["accuracy_metrics_computed"] is False
    assert result.metadata["ausrap"]["sites"][0]["matched_source_rows"] == [2, 3]
    output = write_prediction_outputs(result, tmp_path / "output")
    saved = pd.read_excel(output["workbook"], sheet_name="Prediction_Output")
    assert saved.existing_irap_cross_recoded_ped_per_h.iloc[0] == 7.5
    manifest = json.loads(output["manifest"].read_text())
    assert manifest["ausrap"]["source"]["sha256"] == small_ausrap.sha256
    assert manifest["prediction_basis"] == "released_model_inference"


def test_missing_reference_source_does_not_block_inference(providers, tmp_path):
    providers[0].drop(index=[1, 2], inplace=True)
    source = pd.DataFrame({"site_id": ["001"], "road_no": ["H001"], "slk_km": [356.36]})
    settings = json.loads((ROOT / "configs/prediction/workflow_v1.json").read_text())
    settings["ausrap_source"] = str(tmp_path / "missing.xlsx")
    config = tmp_path / "config.json"
    config.write_text(json.dumps(settings))
    result = predict_sites(source, ROOT, config_path=config)
    assert result.predictions.status.tolist() == ["predicted"]
    assert result.predictions.comparison_status.tolist() == ["source_unavailable"]
    assert result.predictions.pred_peak_cross_ped_per_h.notna().all()
    assert result.predictions.existing_irap_cross_recoded_ped_per_h.isna().all()


def test_reference_remains_available_when_model_features_fail(providers, small_ausrap):
    providers[0].drop(index=[1, 2], inplace=True)
    providers[0].loc[0, "abs_match_status"] = "request_failed"
    source = pd.DataFrame({"site_id": ["001"], "road_no": ["H001"], "slk_km": [356.36]})
    result = predict_sites(source, ROOT, ausrap_index=small_ausrap)
    assert result.predictions.status.tolist() == ["not_predicted"]
    assert result.predictions.pred_peak_cross_ped_per_h.isna().all()
    assert result.predictions.comparison_status.tolist() == ["matched"]
    assert result.predictions.existing_irap_cross_recoded_ped_per_h.tolist() == [7.5]


def test_reference_survives_batch_with_no_resolved_prediction_locations(providers, small_ausrap, tmp_path):
    locator = Mock(path=None)
    locator.by_road.return_value = []
    source = pd.DataFrame({"site_id": ["001"], "road_no": ["H001"], "slk_km": [356.36]})
    result = predict_sites(source, ROOT, locator=locator, ausrap_index=small_ausrap)
    providers[1].assert_not_called()
    providers[3].assert_not_called()
    assert result.predictions.status.tolist() == ["not_predicted"]
    assert result.predictions.comparison_status.tolist() == ["matched"]
    assert result.predictions.existing_irap_cross_recoded_ped_per_h.tolist() == [7.5]
    assert result.predictions.pred_peak_cross_ped_per_h.isna().all()
    output = write_prediction_outputs(result, tmp_path / "output")
    assert json.loads(output["manifest"].read_text())["ausrap"]["status_counts"] == {"matched": 1}
