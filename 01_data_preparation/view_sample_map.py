#!/usr/bin/env python3
"""Build and open an English interactive map from model_dataset.xlsx.

Run from the project directory:

    python3 01_data_preparation/view_sample_map.py

The command regenerates 06_outputs/internal_sample_map.html, starts a local web
server, opens the default browser, and keeps the server running until Ctrl+C.
Use --write-only to generate the HTML without opening a browser.
"""

from __future__ import annotations

import argparse
import json
import math
import threading
import webbrowser
from functools import partial
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except ModuleNotFoundError as exc:
    raise SystemExit(
        "openpyxl is required. Install the project dependencies with "
        "'python3 -m pip install -r requirements.txt'."
    ) from exc


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = PROJECT_ROOT / "data" / "processed" / "model_dataset.xlsx"
DEFAULT_OUTPUT = PROJECT_ROOT / "06_outputs" / "internal_sample_map.html"
REQUIRED_COLUMNS = {
    "site_id",
    "survey_batch_id",
    "location_id",
    "road_no",
    "slk_km",
    "road_name",
    "survey_location_name",
    "latitude",
    "longitude",
    "obs_peak_cross_ped_per_h",
    "obs_peak_along_ped_per_h",
    "qa_status",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Open an interactive map of the pedestrian survey sites."
    )
    parser.add_argument(
        "--workbook",
        type=Path,
        default=DEFAULT_WORKBOOK,
        help="Path to model_dataset.xlsx.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT,
        help="Generated standalone HTML path.",
    )
    parser.add_argument(
        "--port",
        type=int,
        default=0,
        help="Local server port; 0 selects an available port automatically.",
    )
    parser.add_argument(
        "--no-open",
        action="store_true",
        help="Start the local server without opening a browser tab.",
    )
    parser.add_argument(
        "--write-only",
        action="store_true",
        help="Generate the HTML and exit without starting a server.",
    )
    return parser.parse_args()


def clean_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return ""
    return value


def read_sites(workbook_path: Path) -> list[dict[str, Any]]:
    if not workbook_path.is_file():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if "Model_Data" not in workbook.sheetnames:
            raise ValueError("Workbook must contain a Model_Data worksheet.")

        sheet = workbook["Model_Data"]
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]
        missing = sorted(REQUIRED_COLUMNS.difference(headers))
        if missing:
            raise ValueError(f"Model_Data is missing required columns: {', '.join(missing)}")

        sites: list[dict[str, Any]] = []
        for values in rows:
            row = {header: clean_value(value) for header, value in zip(headers, values)}
            if not row["site_id"]:
                continue
            try:
                row["latitude"] = float(row["latitude"])
                row["longitude"] = float(row["longitude"])
                row["slk_km"] = float(row["slk_km"])
                row["obs_peak_cross_ped_per_h"] = float(row["obs_peak_cross_ped_per_h"])
                row["obs_peak_along_ped_per_h"] = float(row["obs_peak_along_ped_per_h"])
            except (TypeError, ValueError) as exc:
                raise ValueError(f"Invalid numeric value for site {row['site_id']}: {exc}") from exc
            sites.append(row)
    finally:
        workbook.close()

    if not sites:
        raise ValueError("Model_Data contains no survey sites.")

    site_ids = [str(site["site_id"]) for site in sites]
    if len(site_ids) != len(set(site_ids)):
        raise ValueError("Model_Data contains duplicate site_id values.")
    if any(not (-90 <= site["latitude"] <= 90) for site in sites):
        raise ValueError("At least one latitude is outside -90 to 90.")
    if any(not (-180 <= site["longitude"] <= 180) for site in sites):
        raise ValueError("At least one longitude is outside -180 to 180.")

    return sites


HTML_TEMPLATE = r"""<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>WA Pedestrian Survey Sample Explorer</title>
  <style>
    /* Required Leaflet 1.9 layout rules are embedded so the local map does not
       depend on a second external stylesheet request. */
    .leaflet-pane, .leaflet-tile, .leaflet-marker-icon, .leaflet-marker-shadow,
    .leaflet-tile-container, .leaflet-pane > svg, .leaflet-pane > canvas,
    .leaflet-zoom-box, .leaflet-image-layer, .leaflet-layer {
      position: absolute; left: 0; top: 0;
    }
    .leaflet-container { overflow: hidden; outline-offset: 1px; -webkit-tap-highlight-color: transparent; }
    .leaflet-container .leaflet-overlay-pane svg { max-width: none !important; max-height: none !important; }
    .leaflet-container .leaflet-tile-pane img, .leaflet-container .leaflet-tile {
      max-width: none !important; max-height: none !important; width: auto; padding: 0;
    }
    .leaflet-tile { user-select: none; -webkit-user-drag: none; visibility: hidden; }
    .leaflet-tile-loaded { visibility: inherit; }
    .leaflet-pane { z-index: 400; }
    .leaflet-tile-pane { z-index: 200; }
    .leaflet-overlay-pane { z-index: 400; }
    .leaflet-shadow-pane { z-index: 500; }
    .leaflet-marker-pane { z-index: 600; }
    .leaflet-tooltip-pane { z-index: 650; }
    .leaflet-popup-pane { z-index: 700; }
    .leaflet-map-pane canvas { z-index: 100; }
    .leaflet-map-pane svg { z-index: 200; }
    .leaflet-control { position: relative; z-index: 800; pointer-events: auto; }
    .leaflet-top, .leaflet-bottom { position: absolute; z-index: 1000; pointer-events: none; }
    .leaflet-top { top: 0; }
    .leaflet-right { right: 0; }
    .leaflet-bottom { bottom: 0; }
    .leaflet-left { left: 0; }
    .leaflet-control { float: left; clear: both; }
    .leaflet-right .leaflet-control { float: right; }
    .leaflet-top .leaflet-control { margin-top: 10px; }
    .leaflet-bottom .leaflet-control { margin-bottom: 10px; }
    .leaflet-left .leaflet-control { margin-left: 10px; }
    .leaflet-right .leaflet-control { margin-right: 10px; }
    .leaflet-zoom-animated { transform-origin: 0 0; }
    .leaflet-zoom-anim .leaflet-zoom-animated { transition: transform 0.25s cubic-bezier(0,0,0.25,1); }
    .leaflet-zoom-anim .leaflet-zoom-hide { visibility: hidden; }
    .leaflet-interactive { cursor: pointer; }
    .leaflet-grab { cursor: grab; }
    .leaflet-dragging .leaflet-grab { cursor: grabbing; }
    .leaflet-marker-icon, .leaflet-marker-shadow, .leaflet-image-layer,
    .leaflet-pane > svg path, .leaflet-tile-container { pointer-events: none; }
    .leaflet-pane > svg path.leaflet-interactive { pointer-events: auto; }
    .leaflet-container { font: 12px/1.5 "Helvetica Neue", Arial, sans-serif; }
    .leaflet-container a { color: #0078a8; }
    .leaflet-bar { box-shadow: 0 1px 5px rgba(0,0,0,0.45); border-radius: 4px; }
    .leaflet-bar a {
      background: #fff; border-bottom: 1px solid #ccc; color: #111; display: block;
      height: 30px; line-height: 30px; text-align: center; text-decoration: none; width: 30px;
    }
    .leaflet-bar a:first-child { border-radius: 4px 4px 0 0; }
    .leaflet-bar a:last-child { border-bottom: 0; border-radius: 0 0 4px 4px; }
    .leaflet-control-zoom-in, .leaflet-control-zoom-out { font: bold 20px "Lucida Console", monospace; }
    .leaflet-container .leaflet-control-attribution {
      background: rgba(255,255,255,0.85); color: #333; line-height: 1.4; margin: 0; padding: 0 5px;
    }
    .leaflet-control-attribution a { text-decoration: none; }
    .leaflet-popup { position: absolute; text-align: center; margin-bottom: 20px; }
    .leaflet-popup-content-wrapper {
      background: #fff; color: #333; padding: 1px; text-align: left;
      border-radius: 9px; box-shadow: 0 3px 14px rgba(0,0,0,0.35);
    }
    .leaflet-popup-content { margin: 13px 24px 13px 20px; line-height: 1.3; min-height: 1px; }
    .leaflet-popup-tip-container {
      width: 40px; height: 20px; position: absolute; left: 50%; margin-left: -20px;
      margin-top: -1px; overflow: hidden; pointer-events: none;
    }
    .leaflet-popup-tip {
      width: 17px; height: 17px; margin: -10px auto 0; padding: 1px;
      background: #fff; box-shadow: 0 3px 14px rgba(0,0,0,0.35); transform: rotate(45deg);
    }
    .leaflet-container a.leaflet-popup-close-button {
      position: absolute; top: 0; right: 0; width: 24px; height: 24px; border: 0;
      background: transparent; color: #666; font: 16px/24px Tahoma, sans-serif;
      text-align: center; text-decoration: none;
    }
    .leaflet-tooltip {
      position: absolute; padding: 6px; background: #fff; border: 1px solid #fff;
      border-radius: 3px; color: #222; white-space: nowrap; pointer-events: none;
      box-shadow: 0 1px 3px rgba(0,0,0,0.35);
    }
    .leaflet-tooltip-top { margin-top: -6px; }
    :root {
      color-scheme: light;
      font-family: Inter, ui-sans-serif, system-ui, -apple-system, BlinkMacSystemFont, "Segoe UI", sans-serif;
      color: #172033;
      background: #f5f7fa;
    }
    * { box-sizing: border-box; }
    body { margin: 0; min-width: 320px; }
    header {
      padding: 16px 20px 14px;
      background: #ffffff;
      border-bottom: 1px solid #d8dee8;
    }
    h1 { margin: 0 0 4px; font-size: 22px; font-weight: 650; }
    header p { margin: 0; color: #5d687b; font-size: 14px; }
    main {
      display: grid;
      grid-template-columns: minmax(270px, 330px) minmax(0, 1fr);
      min-height: calc(100vh - 80px);
    }
    aside {
      padding: 18px;
      background: #ffffff;
      border-right: 1px solid #d8dee8;
      overflow-y: auto;
    }
    label { display: block; margin: 0 0 6px; font-size: 13px; font-weight: 650; }
    select, button {
      width: 100%;
      min-height: 40px;
      border: 1px solid #b8c1cf;
      border-radius: 6px;
      background: #ffffff;
      color: #172033;
      font: inherit;
    }
    select { margin-bottom: 16px; padding: 8px 10px; }
    button { padding: 8px 12px; cursor: pointer; font-weight: 600; }
    button:hover { background: #f0f3f7; }
    .detail {
      margin-top: 18px;
      padding-top: 16px;
      border-top: 1px solid #d8dee8;
    }
    .detail h2 { margin: 0 0 10px; font-size: 16px; }
    .detail dl { margin: 0; display: grid; grid-template-columns: 92px 1fr; gap: 7px 10px; font-size: 13px; }
    .detail dt { color: #667085; }
    .detail dd { margin: 0; overflow-wrap: anywhere; }
    .hint { margin: 12px 0 0; color: #667085; font-size: 12px; line-height: 1.45; }
    #map { min-height: calc(100vh - 80px); background: #e8edf2; }
    .map-legend {
      padding: 8px 10px;
      background: rgba(255, 255, 255, 0.96);
      border: 1px solid #cbd3df;
      border-radius: 5px;
      box-shadow: 0 1px 4px rgba(23, 32, 51, 0.15);
      color: #172033;
      font-size: 12px;
      line-height: 1.45;
    }
    .legend-row { display: flex; align-items: center; gap: 7px; margin-top: 4px; }
    .legend-dot { width: 11px; height: 11px; border-radius: 50%; display: inline-block; flex: none; }
    .popup-table { border-collapse: collapse; font-size: 12px; min-width: 250px; }
    .popup-table th, .popup-table td { padding: 3px 5px; text-align: left; vertical-align: top; }
    .popup-table th { color: #5d687b; font-weight: 600; white-space: nowrap; }
    @media (max-width: 720px) {
      main { grid-template-columns: 1fr; }
      aside { border-right: 0; border-bottom: 1px solid #d8dee8; overflow: visible; }
      #map { min-height: 65vh; }
    }
  </style>
</head>
<body>
  <header>
    <h1>WA Pedestrian Survey Sample Explorer</h1>
    <p>__SITE_COUNT__ surveyed locations · Source: model_dataset.xlsx · Select a site to zoom to its city and road context.</p>
  </header>
  <main>
    <aside>
      <label for="display-mode">Display markers by</label>
      <select id="display-mode">
        <option value="batch">Survey batch</option>
        <option value="cross">Peak cross-road pedestrian flow</option>
        <option value="along">Peak along-road pedestrian flow</option>
      </select>

      <label for="site-select">Find a site, locality or road</label>
      <select id="site-select">
        <option value="">All surveyed locations</option>
      </select>
      <button id="reset-view" type="button">Show all Western Australia sites</button>

      <section class="detail" aria-live="polite">
        <h2>Selected location</h2>
        <div id="site-detail">Choose a location from the list or click a marker.</div>
      </section>
      <p class="hint">Zoom in to read city, suburb and road labels from the OpenStreetMap basemap. Marker tooltips show the survey location label.</p>
    </aside>
    <div id="map" role="application" aria-label="Interactive map of Western Australian pedestrian survey locations"></div>
  </main>

  <script src="https://unpkg.com/leaflet@1.9.4/dist/leaflet.js"
          integrity="sha256-20nQCchB9co0qIjJZRGuk2/Z9VM+kNiyxNV1lvTlZBo=" crossorigin=""></script>
  <script>
    const sites = __POINTS_JSON__;
    const batchPalette = ["#2563eb", "#f97316", "#16a34a", "#9333ea", "#dc2626", "#0891b2"];
    const batchIds = [...new Set(sites.map(d => String(d.survey_batch_id)))].sort();
    const batchColors = Object.fromEntries(batchIds.map((id, index) => [id, batchPalette[index % batchPalette.length]]));
    const modeSelect = document.getElementById("display-mode");
    const siteSelect = document.getElementById("site-select");
    const detail = document.getElementById("site-detail");

    const map = L.map("map", { zoomControl: true, preferCanvas: true });
    L.tileLayer("https://{s}.tile.openstreetmap.org/{z}/{x}/{y}.png", {
      maxZoom: 19,
      attribution: "&copy; OpenStreetMap contributors"
    }).addTo(map);

    const bounds = L.latLngBounds(sites.map(d => [d.latitude, d.longitude]));
    const markerById = new Map();
    let selectedSiteId = "";

    function esc(value) {
      return String(value ?? "").replace(/[&<>'"]/g, ch => ({
        "&": "&amp;", "<": "&lt;", ">": "&gt;", "'": "&#39;", '"': "&quot;"
      })[ch]);
    }

    function number(value, digits = 1) {
      return Number(value).toLocaleString("en-AU", { maximumFractionDigits: digits });
    }

    function markerRadius(site, mode) {
      if (mode === "batch") return 6;
      const key = mode === "cross" ? "obs_peak_cross_ped_per_h" : "obs_peak_along_ped_per_h";
      const maxValue = Math.max(...sites.map(d => Number(d[key])), 1);
      return 4 + 13 * Math.sqrt(Math.max(0, Number(site[key])) / maxValue);
    }

    function markerColor(site, mode) {
      if (mode === "batch") return batchColors[String(site.survey_batch_id)];
      return mode === "cross" ? "#2563eb" : "#f97316";
    }

    function popupHtml(site) {
      return `<table class="popup-table">
        <tr><th>Site ID</th><td>${esc(site.site_id)}</td></tr>
        <tr><th>Locality / label</th><td>${esc(site.survey_location_name)}</td></tr>
        <tr><th>Road</th><td>${esc(site.road_no)} · ${esc(site.road_name)}</td></tr>
        <tr><th>SLK</th><td>${number(site.slk_km, 3)} km</td></tr>
        <tr><th>Coordinates</th><td>${number(site.latitude, 6)}, ${number(site.longitude, 6)}</td></tr>
        <tr><th>Survey batch</th><td>${esc(site.survey_batch_id)}</td></tr>
        <tr><th>Cross-road peak</th><td>${number(site.obs_peak_cross_ped_per_h, 3)} pedestrians/hour</td></tr>
        <tr><th>Along-road peak</th><td>${number(site.obs_peak_along_ped_per_h, 3)} pedestrians/hour</td></tr>
        <tr><th>QA status</th><td>${esc(site.qa_status)}</td></tr>
      </table>`;
    }

    function updateDetail(site) {
      if (!site) {
        detail.textContent = "Choose a location from the list or click a marker.";
        return;
      }
      detail.innerHTML = `<dl>
        <dt>Site ID</dt><dd>${esc(site.site_id)}</dd>
        <dt>Locality</dt><dd>${esc(site.survey_location_name)}</dd>
        <dt>Road</dt><dd>${esc(site.road_no)} · ${esc(site.road_name)}</dd>
        <dt>SLK</dt><dd>${number(site.slk_km, 3)} km</dd>
        <dt>Cross-road</dt><dd>${number(site.obs_peak_cross_ped_per_h, 3)} pedestrians/hour</dd>
        <dt>Along-road</dt><dd>${number(site.obs_peak_along_ped_per_h, 3)} pedestrians/hour</dd>
      </dl>`;
    }

    function selectSite(site, zoom = true) {
      selectedSiteId = site.site_id;
      siteSelect.value = site.site_id;
      updateDetail(site);
      updateMarkers();
      const marker = markerById.get(site.site_id);
      if (zoom) map.setView(marker.getLatLng(), 14, { animate: true });
      marker.openPopup();
    }

    sites
      .slice()
      .sort((a, b) => `${a.survey_location_name} ${a.site_id}`.localeCompare(`${b.survey_location_name} ${b.site_id}`, "en"))
      .forEach(site => {
        const option = document.createElement("option");
        option.value = site.site_id;
        option.textContent = `${site.site_id} — ${site.survey_location_name} — ${site.road_name}`;
        siteSelect.appendChild(option);

        const marker = L.circleMarker([site.latitude, site.longitude], {
          radius: 6,
          color: "#ffffff",
          weight: 1.5,
          fillColor: batchColors[String(site.survey_batch_id)],
          fillOpacity: 0.86
        }).addTo(map);
        marker.bindTooltip(`${esc(site.site_id)} · ${esc(site.survey_location_name)} · ${esc(site.road_name)}`, { direction: "top" });
        marker.bindPopup(popupHtml(site), { maxWidth: 360 });
        marker.on("click", () => selectSite(site, false));
        markerById.set(site.site_id, marker);
      });

    const legend = L.control({ position: "bottomright" });
    legend.onAdd = function () {
      this._div = L.DomUtil.create("div", "map-legend");
      this.update();
      return this._div;
    };
    legend.update = function () {
      const mode = modeSelect.value;
      if (mode === "batch") {
        this._div.innerHTML = `<strong>Survey batch</strong>${batchIds.map(id =>
          `<div class="legend-row"><span class="legend-dot" style="background:${batchColors[id]}"></span>${esc(id)}</div>`
        ).join("")}`;
      } else {
        const label = mode === "cross" ? "Cross-road peak" : "Along-road peak";
        const color = mode === "cross" ? "#2563eb" : "#f97316";
        this._div.innerHTML = `<strong>${label}</strong><div class="legend-row"><span class="legend-dot" style="background:${color}"></span>Larger circle = more pedestrians/hour</div>`;
      }
    };
    legend.addTo(map);

    function updateMarkers() {
      const mode = modeSelect.value;
      sites.forEach(site => {
        const marker = markerById.get(site.site_id);
        const selected = site.site_id === selectedSiteId;
        marker.setRadius(markerRadius(site, mode) + (selected ? 2 : 0));
        marker.setStyle({
          fillColor: markerColor(site, mode),
          color: selected ? "#111827" : "#ffffff",
          weight: selected ? 3 : 1.5,
          fillOpacity: 0.86
        });
      });
      legend.update();
    }

    modeSelect.addEventListener("change", updateMarkers);
    siteSelect.addEventListener("change", () => {
      const site = sites.find(d => d.site_id === siteSelect.value);
      if (site) selectSite(site, true);
      else resetView();
    });

    function resetView() {
      selectedSiteId = "";
      siteSelect.value = "";
      updateDetail(null);
      updateMarkers();
      map.closePopup();
      map.fitBounds(bounds, { padding: [24, 24] });
    }

    document.getElementById("reset-view").addEventListener("click", resetView);
    map.fitBounds(bounds, { padding: [24, 24] });
    updateMarkers();
  </script>
</body>
</html>
"""


def build_html(sites: list[dict[str, Any]], workbook_path: Path) -> str:
    points_json = json.dumps(sites, ensure_ascii=False, separators=(",", ":")).replace(
        "</", "<\\/"
    )
    return (
        HTML_TEMPLATE.replace("__SITE_COUNT__", str(len(sites)))
        .replace("__POINTS_JSON__", points_json)
        .replace("model_dataset.xlsx", workbook_path.name)
    )


class QuietHandler(SimpleHTTPRequestHandler):
    def log_message(self, format: str, *args: object) -> None:
        return


def serve(output_path: Path, port: int, open_browser: bool) -> None:
    handler = partial(QuietHandler, directory=str(output_path.parent))
    server = ThreadingHTTPServer(("127.0.0.1", port), handler)
    url = f"http://127.0.0.1:{server.server_port}/{output_path.name}"
    print(f"Map ready: {url}")
    print("Press Ctrl+C in this terminal to stop the local server.")
    if open_browser:
        threading.Timer(0.4, lambda: webbrowser.open_new_tab(url)).start()
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nLocal map server stopped.")
    finally:
        server.server_close()


def main() -> None:
    args = parse_args()
    workbook_path = args.workbook.expanduser().resolve()
    output_path = args.output.expanduser().resolve()
    sites = read_sites(workbook_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(build_html(sites, workbook_path), encoding="utf-8")
    print(f"Generated {output_path} from {len(sites)} unique survey sites.")
    if args.write_only:
        return
    serve(output_path, args.port, open_browser=not args.no_open)


if __name__ == "__main__":
    main()
