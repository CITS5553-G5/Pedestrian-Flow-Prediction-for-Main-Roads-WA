"""Fill the maintained formal prediction template and write run provenance."""

from copy import copy
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment

from pedflow.exceptions import ContractError
from pedflow.io import atomic_path, write_json_atomic
from pedflow.prediction.schema import FEATURE_COLUMNS, OUTPUT_COLUMNS
from pedflow.prediction.workflow import PredictionResult


def metadata_table(metadata: dict) -> pd.DataFrame:
    """Readable run facts; RAP benchmark data never enters formal inference."""
    fields = (
        "run_id", "output_schema_version", "model_version", "coordinate_crs",
        "prediction_units", "started_at_utc", "finished_at_utc", "row_count",
        "road_snapshot_sha256",
    )
    values = {name: metadata.get(name, "") for name in fields}
    values.update(
        selected_features="commercial_buildings, footpath_both_sides, abs_mb_category_simple",
        location_rule="Road + SLK uses measured road geometry; coordinates use the entered point. Both inputs use coordinates for feature extraction and check the road reference.",
        missing_values="Blank means unavailable, not zero. Location and feature statuses explain failures.",
    )
    for status, count in metadata.get("location_status_counts", {}).items():
        values[f"location: {status}"] = count
    return pd.DataFrame(values.items(), columns=["field", "value"])


def write_prediction_outputs(
    result: PredictionResult,
    directory: str | Path,
    *,
    template_path: str | Path | None = None,
) -> dict[str, Path]:
    """Fill the existing styled output template for this run only."""
    folder = Path(directory)
    workbook = folder / "predictions.xlsx"
    template = Path(template_path) if template_path else (
        Path(__file__).resolve().parents[3] / "data/templates/prediction_output_template.xlsx"
    )
    book = load_workbook(template)
    for sheet_name, table, schema in (
        ("Prediction_Output", result.predictions, OUTPUT_COLUMNS),
        ("Features", result.features, FEATURE_COLUMNS),
    ):
        sheet = book[sheet_name]
        headers = tuple(sheet.cell(1, col).value for col in range(1, len(schema) + 1))
        if headers != schema:
            raise ContractError(f"Output template {sheet_name} headers do not match the service schema.")
        if table.site_id.isna().any() or table.site_id.astype("string").str.strip().eq("").any() or table.site_id.duplicated().any():
            raise ContractError(f"{sheet_name} requires unique, nonblank site_id values.")
        for row in sheet.iter_rows(min_row=2, max_col=len(schema)):
            for cell in row:
                cell.value = None
        data = table.reindex(columns=schema).astype(object)
        data = data.where(pd.notna(data), None)
        for row_number, values in enumerate(data.itertuples(index=False, name=None), start=2):
            for col_number, value in enumerate(values, start=1):
                cell = sheet.cell(row_number, col_number)
                if row_number > 2 and not cell.has_style:
                    cell._style = copy(sheet.cell(2, col_number)._style)
                cell.value = value.item() if isinstance(value, np.generic) else value
                if col_number == 1 and isinstance(value, str) and value.startswith("="):
                    cell.data_type = "s"
                heading = schema[col_number - 1]
                if heading in {"latitude", "longitude"}:
                    cell.number_format = "0.000000"
                elif heading in {"slk_km", "start_slk", "end_slk", "matched_slk_km"}:
                    cell.number_format = "0.000"
                elif heading.endswith("_per_h"):
                    cell.number_format = "0.00"
            sheet.row_dimensions[row_number].height = 16
        sheet.freeze_panes = "B2"
        last_column = sheet.cell(1, len(schema)).column_letter
        sheet.auto_filter.ref = f"A1:{last_column}{max(2, len(data) + 1)}"
        sheet.sheet_view.showGridLines = False
    run_sheet = book["Metadata"]
    run_sheet.cell(50, 1, "Run metadata")
    run_sheet.cell(51, 1, "field")
    run_sheet.cell(51, 2, "value")
    for row_number, (field, value) in enumerate(metadata_table(result.metadata).itertuples(index=False, name=None), start=52):
        run_sheet.cell(row_number, 1, field)
        run_sheet.cell(row_number, 2, value)
        run_sheet.cell(row_number, 2).alignment = Alignment(wrap_text=True, vertical="top")
        run_sheet.row_dimensions[row_number].height = 42 if isinstance(value, str) and len(value) > 80 else 16
    with atomic_path(workbook) as temporary:
        book.save(temporary)
    manifest = folder / "predictions.manifest.json"
    write_json_atomic(manifest, result.metadata)
    return {"workbook": workbook, "manifest": manifest}
