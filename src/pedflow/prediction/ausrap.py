"""Local AusRAP reference matching, separate from model inference and accuracy metrics."""

import hashlib
import json
import math
import re
import sqlite3
from collections import Counter
from contextlib import closing
from pathlib import Path
from zipfile import BadZipFile

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from pedflow.exceptions import ContractError
from pedflow.flow import BANDS, CATEGORY_RULE, RECOUNT, WORD_CODES, flow_category, format_band
from pedflow.io import atomic_path

DEFAULT_SOURCE = "data/raw/AusRAP Data - Words - MSDS.xlsx"
MATCH_VERSION = "ausrap_reference_v1"
INDEX_VERSION = "ausrap_index_v1"
WINDOW_KM = 0.1
SOURCE_COLUMNS = (
    "Road No", "Start SLK", "End SLK", "CWAY",
    "Pedestrian peak hour flow across the road",
    "Pedestrian peak hour flow along the road driver-side",
    "Pedestrian peak hour flow along the road passenger-side",
)
REFERENCE_COLUMNS = (
    "existing_irap_cross_recoded_ped_per_h", "existing_ausrap_cross_category",
    "existing_ausrap_cross_band", "existing_irap_along_recoded_ped_per_h",
    "existing_ausrap_along_combined_band", "comparison_status", "comparison_notes",
)


def _text(value):
    return "" if pd.isna(value) else str(value).strip()


def _number(value):
    if isinstance(value, bool):
        return None
    try:
        value = float(value)
        return value if math.isfinite(value) else None
    except (ValueError, TypeError):
        return None


def _code(value):
    if _text(value) in WORD_CODES:
        return WORD_CODES[_text(value)]
    # Also accept explicit category codes; never interpret an unknown label as a count.
    number = _number(value)
    return int(number) if number is not None and number.is_integer() and 1 <= number <= 11 else None


def _sha256(path):
    with Path(path).open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


class AusRAPIndex:
    """Content-addressed SQLite cache of the original source rows; no deduplication."""

    def __init__(self, source, cache_directory):
        self.source = Path(source).resolve()
        self.sha256 = _sha256(self.source)
        self.path = Path(cache_directory) / f"{INDEX_VERSION}-{self.sha256}.sqlite"
        if not self.path.exists():
            self._build()
        with self._connect() as connection:
            saved = connection.execute("SELECT payload FROM metadata").fetchone()
            if saved is None:
                raise ContractError("AusRAP cache metadata is missing")
            self.metadata = json.loads(saved[0])
        if (not isinstance(self.metadata, dict) or self.metadata.get("sha256") != self.sha256
                or self.metadata.get("index_version") != INDEX_VERSION):
            raise ContractError("AusRAP cache does not match its source/version")
        self.metadata["source_path"] = str(self.source)

    def _connect(self):
        return closing(sqlite3.connect(f"{self.path.resolve().as_uri()}?mode=ro", uri=True))

    def _build(self):
        self.path.parent.mkdir(parents=True, exist_ok=True)
        book = load_workbook(self.source, read_only=True, data_only=True)
        try:
            sheet = book.worksheets[0]
            headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
            if any(headers.count(name) != 1 for name in SOURCE_COLUMNS):
                raise ContractError("AusRAP source is missing unique required road/flow headers")
            positions = [headers.index(name) for name in SOURCE_COLUMNS]
            counts = Counter()
            with atomic_path(self.path) as temporary, closing(sqlite3.connect(temporary)) as db:
                db.execute("CREATE TABLE segments (source_row INTEGER PRIMARY KEY, road TEXT, "
                           "start REAL, end REAL, cway TEXT, cross_code INTEGER, driver_code INTEGER, "
                           "passenger_code INTEGER, cross_label TEXT, driver_label TEXT, passenger_label TEXT)")
                pending = []
                for source_row, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
                    road, start, end, cway, *labels = [row[p] for p in positions]
                    start, end = _number(start), _number(end)
                    road, cway = _text(road).upper(), _text(cway).upper()
                    if start is None or end is None or not road:
                        counts["invalid_location_rows"] += 1
                    elif end < start:
                        start, end = end, start
                        counts["reversed_bounds_rows"] += 1
                    codes = [_code(label) for label in labels]
                    for label, code in zip(labels, codes):
                        if code is None:
                            counts["missing_flow_cells" if not _text(label) else "unrecognized_flow_cells"] += 1
                    pending.append((source_row, road, start, end, cway, *codes,
                                    *[_text(label) or None for label in labels]))
                    counts["source_rows"] += 1
                    if len(pending) == 1000:
                        db.executemany("INSERT INTO segments VALUES (?,?,?,?,?,?,?,?,?,?,?)", pending)
                        pending.clear()
                db.executemany("INSERT INTO segments VALUES (?,?,?,?,?,?,?,?,?,?,?)", pending)
                if not counts["source_rows"]:
                    raise ContractError("AusRAP source has no records")
                db.execute("CREATE INDEX road_bounds ON segments(road,start,end)")
                metadata = {"source_path": str(self.source), "sha256": self.sha256,
                            "source_sheet": sheet.title, "index_version": INDEX_VERSION, **counts}
                if _sha256(self.source) != self.sha256:
                    raise ContractError("AusRAP source changed during indexing; retry with a stable file")
                db.execute("CREATE TABLE metadata (payload TEXT)")
                db.execute("INSERT INTO metadata VALUES (?)", (json.dumps(metadata),))
                db.commit()
        finally:
            book.close()

    def query(self, road, lower, upper):
        """Client overlap convention: direct floats, inclusive, all carriageways."""
        with self._connect() as db:
            db.row_factory = sqlite3.Row
            return [dict(row) for row in db.execute(
                "SELECT * FROM segments WHERE road=? AND start<=? AND end>=? ORDER BY source_row",
                (str(road).strip().upper(), upper, lower),
            )]

    def record(self, record_id):
        """IDs are source SHA-256 plus Excel row, not purported official identifiers."""
        match = re.fullmatch(r"ausrap:([a-f0-9]{64}):([0-9]+)", str(record_id).strip())
        if not match or match[1] != self.sha256 or len(match[2]) > 7:
            raise ContractError("AusRAP record ID is invalid or belongs to a different source snapshot")
        with self._connect() as db:
            db.row_factory = sqlite3.Row
            row = db.execute("SELECT * FROM segments WHERE source_row=?", (int(match[2]),)).fetchone()
        if row is None:
            raise ContractError("AusRAP record ID does not exist in this source")
        return dict(row)


def _aggregate(records):
    result = dict.fromkeys(REFERENCE_COLUMNS)
    notes = []
    cross = [row["cross_code"] for row in records]
    along = [(row["driver_code"], row["passenger_code"]) for row in records]
    valid_cross = all(code in BANDS for code in cross)
    valid_along = all(a in BANDS and b in BANDS for a, b in along)
    code_set = sorted({code for code in cross if code is not None})
    if valid_cross and len(code_set) == 1:
        result["existing_ausrap_cross_category"] = code_set[0]
    if len(code_set) > 1:
        notes.append("Mixed crossing categories: " + ", ".join(map(str, code_set)))

    def average_band(pairs):
        return (sum(p[0] for p in pairs) / len(pairs),
                None if any(p[1] is None for p in pairs) else sum(p[1] for p in pairs) / len(pairs))

    cross_band = average_band([BANDS[c] for c in cross]) if valid_cross else None
    along_band = None
    if valid_along:
        combined = [(BANDS[a][0] + BANDS[b][0],
                     None if BANDS[a][1] is None or BANDS[b][1] is None else BANDS[a][1] + BANDS[b][1])
                    for a, b in along]
        along_band = average_band(combined)
    result["existing_ausrap_cross_band"] = format_band(cross_band)
    result["existing_ausrap_along_combined_band"] = format_band(along_band)
    if all(c in RECOUNT for c in cross):
        result["existing_irap_cross_recoded_ped_per_h"] = sum(RECOUNT[c] for c in cross) / len(cross)
    if all(a in RECOUNT and b in RECOUNT for a, b in along):
        result["existing_irap_along_recoded_ped_per_h"] = sum(RECOUNT[a] + RECOUNT[b] for a, b in along) / len(along)
    if not valid_cross:
        notes.append("Crossing source category missing or unrecognized; full matched set retained")
    if not valid_along:
        notes.append("Along-road source category missing or unrecognized; full matched set retained")
    if any(c in (10, 11) for c in cross + [c for pair in along for c in pair]):
        notes.append("Client recoding is undefined for category 10/11; affected numeric estimate is blank")
    cways = sorted({_text(row["cway"]) for row in records})
    if len(cways) > 1:
        notes.append("Client mean spans multiple carriageways")
    if len(records) > 1:
        notes.append("Intervals are means of source endpoints; along-road combines both sides first")
    numeric = [result[k] is not None for k in (
        "existing_irap_cross_recoded_ped_per_h", "existing_irap_along_recoded_ped_per_h")]
    result["comparison_status"] = ("matched" if all(numeric) else "matched_partial" if any(numeric)
                                   else "matched_bands_only" if valid_cross or valid_along else "matched_no_flow")
    result["comparison_notes"] = "; ".join(notes)
    return result, {"cross_categories": code_set, "cross_band": cross_band,
                    "along_combined_band": along_band, "carriageways": cways}


def match_site(location, index):
    """Match a resolved location without changing any model result."""
    result = dict.fromkeys(REFERENCE_COLUMNS)
    audit = {"site_id": str(location["site_id"]), "matched_source_rows": [], "method": None}

    def unavailable(status, note):
        result.update(comparison_status=status, comparison_notes=note)
        audit["status"] = status
        return result, audit

    if location.get("location_status") == "road_coordinate_conflict":
        return unavailable("location_conflict", "Road reference conflicts with prediction coordinates")
    road = _text(location.get("road_no"))
    if road:
        slk = _number(location.get("slk_km"))
        if slk is None:
            start, end = _number(location.get("start_slk")), _number(location.get("end_slk"))
            slk = (start + end) / 2 if start is not None and end is not None else None
        audit["location_basis"] = "supplied_road_reference"
    else:
        if location.get("location_status") != "matched":
            return unavailable("location_unresolved", "No unambiguous road/SLK reference for AusRAP matching")
        road = _text(location.get("matched_road_no"))
        slk = _number(location.get("_matched_slk_km_exact", location.get("matched_slk_km")))
        audit["location_basis"] = "coordinate_road_match"
    if not road or slk is None:
        return unavailable("location_unresolved", "Road number or SLK unavailable for AusRAP matching")
    road = road.upper()
    audit.update(road_no=road, slk_km=slk)
    record_id = _text(location.get("ausrap_record_id"))
    if record_id:
        audit["method"] = "source_record_id"
        try:
            record = index.record(record_id)
        except ContractError as exc:
            return unavailable("invalid_record_id", str(exc))
        if (record["road"] != road or record["start"] is None or record["end"] is None
                or not record["start"] <= slk <= record["end"]
                or (_text(location.get("cway")) and _text(location.get("cway")).upper() != record["cway"])):
            return unavailable("record_location_conflict", "AusRAP record ID conflicts with supplied road/SLK/carriageway")
        records = [record]
    else:
        audit.update(method="point_window_equal_mean", window_km=[slk - WINDOW_KM, slk + WINDOW_KM])
        records = index.query(road, slk - WINDOW_KM, slk + WINDOW_KM)
    if not records:
        return unavailable("no_ausrap_match", "No AusRAP source rows overlap the road/SLK window")
    result, detail = _aggregate(records)
    input_cway = _text(location.get("cway")).upper()
    if not record_id and input_cway and any(r["cway"] != input_cway for r in records):
        note = "Point-window reference includes carriageways other than the input; client rule does not filter carriageways"
        result["comparison_notes"] = "; ".join(filter(None, [result["comparison_notes"], note]))
    if location.get("location_status") != "matched":
        note = "Reference uses supplied road/SLK; prediction location status: " + _text(location.get("location_status"))
        result["comparison_notes"] = "; ".join(filter(None, [result["comparison_notes"], note]))
    audit.update(detail, status=result["comparison_status"], matched_source_rows=[r["source_row"] for r in records],
                 matched_record_ids=[f"ausrap:{index.sha256}:{r['source_row']}" for r in records])
    return result, audit


def attach_reference(predictions, locations, root, source_path=DEFAULT_SOURCE, *, index=None):
    """Append presentation-only references. Missing references never erase predictions."""
    table = predictions.copy(deep=True)
    indexed_locations = locations.set_index("site_id", drop=False)
    metadata = {
        "version": MATCH_VERSION, "category_rule": CATEGORY_RULE,
        "reference_label": "Existing iRAP flow estimate (recoded)",
        "prediction_basis": "released_model_inference", "accuracy_metrics_computed": False,
        "point_window_half_width_km": WINDOW_KM, "carriageway_filter_for_point_window": False,
        "aggregation": "Equal-weight mean of all matched source rows; no deduplication or missing-value deletion.",
        "interval_meaning": "Derived source endpoint means, not model confidence intervals.",
        "record_id_format": "ausrap:<source SHA-256>:<Excel row number>",
        "source": None, "sites": [],
    }
    source_error = None
    if index is None:
        source = Path(source_path)
        source = source if source.is_absolute() else Path(root) / source
        try:
            index = AusRAPIndex(source, Path(root) / "data/cache/ausrap")
        except (OSError, ValueError, ContractError, sqlite3.Error, BadZipFile, InvalidFileException) as exc:
            source_error = str(exc)
            metadata.update(source_path=str(source), source_error=source_error)
    if index is not None:
        metadata["source"] = index.metadata
    appended = []
    for row in table.to_dict("records"):
        location = indexed_locations.loc[row["site_id"]].to_dict()
        if source_error is None:
            reference, audit = match_site(location, index)
        else:
            reference = dict.fromkeys(REFERENCE_COLUMNS)
            reference.update(comparison_status="source_unavailable", comparison_notes="AusRAP source unavailable: " + source_error)
            audit = {"site_id": str(row["site_id"]), "status": "source_unavailable", "matched_source_rows": []}
        row.update(reference)
        for direction in ("cross", "along"):
            row[f"obs_peak_{direction}_ped_per_h"] = _number(location.get(f"obs_peak_{direction}_ped_per_h"))
        row["observed_cross_category"] = flow_category(row["obs_peak_cross_ped_per_h"])
        row["predicted_cross_category"] = flow_category(row.get("pred_peak_cross_ped_per_h"))
        appended.append(row)
        metadata["sites"].append(audit)
    metadata["status_counts"] = dict(Counter(row["comparison_status"] for row in appended))
    table = pd.DataFrame(appended)
    for name in ("observed_cross_category", "predicted_cross_category", "existing_ausrap_cross_category"):
        table[name] = table[name].astype("Int64")
    return table, metadata
