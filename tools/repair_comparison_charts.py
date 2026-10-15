"""Repair missing XY references in artifact-tool's native scatter chart export.

Only the two chart XML parts are changed. Worksheets, formula results and all
other workbook parts remain byte-identical. openpyxl is used only for reading.
"""

import os
import sys
import xml.etree.ElementTree as etree
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

path = Path(sys.argv[1])
book = load_workbook(path, data_only=True)
namespace = "http://schemas.openxmlformats.org/drawingml/2006/chart"
ns = {"c": namespace}


def numeric_reference(parent, reference):
    for child in list(parent):
        parent.remove(child)
    ref = etree.SubElement(parent, f"{{{namespace}}}numRef")
    etree.SubElement(ref, f"{{{namespace}}}f").text = reference
    cache = etree.SubElement(ref, f"{{{namespace}}}numCache")
    etree.SubElement(cache, f"{{{namespace}}}formatCode").text = "0.00"
    sheet, address = reference.split("!")
    sheet = sheet.strip("'")
    first_col, first_row, last_col, last_row = range_boundaries(address)
    assert first_col == last_col
    values = [book[sheet].cell(r, first_col).value for r in range(first_row, last_row + 1)]
    assert all(isinstance(v, (float, int)) and not isinstance(v, bool) for v in values)
    etree.SubElement(cache, f"{{{namespace}}}ptCount", val=str(len(values)))
    for i, value in enumerate(values):
        point = etree.SubElement(cache, f"{{{namespace}}}pt", idx=str(i))
        etree.SubElement(point, f"{{{namespace}}}v").text = repr(value)


replacements = {}
with ZipFile(path) as archive:
    for chart_number, xcolumn, count in [(1, "O", 45), (2, "T", 43)]:
        name = f"xl/drawings/charts/chart{chart_number}.xml"
        root = etree.fromstring(archive.read(name))
        series = root.findall(".//c:ser", ns)
        assert len(series) == 3
        for item in series:
            x = item.find("c:xVal", ns)
            y = item.find("c:yVal", ns)
            yreference = y.findtext("c:numRef/c:f", namespaces=ns)
            assert yreference and yreference.startswith("'Errors'!")
            numeric_reference(x, f"'Errors'!${xcolumn}$2:${xcolumn}${count+1}")
            numeric_reference(y, yreference)
        replacements[name] = etree.tostring(root, encoding="UTF-8", xml_declaration=True)
    temporary = path.with_suffix(".xy-repaired.xlsx")
    with ZipFile(temporary, "w") as target:
        for item in archive.infolist():
            target.writestr(item, replacements.get(item.filename, archive.read(item.filename)))
with ZipFile(path) as old, ZipFile(temporary) as new:
    changed = [name for name in old.namelist() if old.read(name) != new.read(name)]
    assert set(changed).issubset(replacements), changed
os.replace(temporary, path)
print("Repaired only the two native scatter chart XML parts; XY references and caches populated.")
