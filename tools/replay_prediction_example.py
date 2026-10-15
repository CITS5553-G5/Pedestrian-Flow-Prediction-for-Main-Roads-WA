"""Replay the delivered feature snapshot without external service requests.

Run from the repository root after installing the project, or with PYTHONPATH=src.
This is a reproducible demonstration of released-model inference and reference
attachment; use pedflow-predict for new locations and fresh feature extraction.
"""

import argparse
import hashlib
import json
from datetime import UTC, datetime
from pathlib import Path
from uuid import uuid4

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from pedflow.prediction.ausrap import attach_reference
from pedflow.prediction.config import load_settings
from pedflow.prediction.io import write_prediction_outputs
from pedflow.prediction.location import RoadLocator, resolve_locations
from pedflow.prediction.schema import OUTPUT_COLUMNS, OUTPUT_SCHEMA_VERSION
from pedflow.prediction.workflow import PREDICTION_COLUMNS, PredictionResult

ROOT = Path(__file__).resolve().parents[1]
SNAPSHOT = "04_prediction_pipeline/prediction_example/evidence/saved_run.json"
INPUT_COLUMNS = ["site_id", "road_no", "slk_km", "cway", "start_slk", "end_slk",
                 "latitude", "longitude"]


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-directory", required=True, type=Path)
    args = parser.parse_args()
    output = args.output_directory
    if output.exists() and any(output.iterdir()):
        parser.error("Use a new empty output directory; reviewed results are not overwritten.")
    source = ROOT / SNAPSHOT
    source_sha = sha(source)
    payload = json.loads(source.read_text())
    if payload["snapshot_version"] != 1:
        parser.error("Unsupported saved-feature snapshot version.")
    started = datetime.now(UTC).isoformat()
    old = pd.DataFrame(**payload["predictions"], dtype=object)
    features = pd.DataFrame(**payload["features"], dtype=object)
    original_metadata = payload["metadata"]
    settings, _, _, model = load_settings(ROOT, "configs/prediction/workflow_v1.json")
    for field in ("model_release_sha256", "imagery_recipe_sha256", "context_feature_set_version"):
        if settings[field] != original_metadata["configuration"][field]:
            raise ValueError(f"Saved snapshot does not match configured {field}")
    locator = RoadLocator(ROOT / "data/external_raw/main_roads_slk/road_slk.sqlite.gz")
    locations = resolve_locations(old[INPUT_COLUMNS], locator=locator)
    assert locations.site_id.tolist() == old.site_id.tolist() == features.site_id.tolist()
    assert locations.location_status.eq("matched").all()
    np.testing.assert_allclose(locations[["latitude", "longitude"]].astype(float),
                               old[["latitude", "longitude"]].astype(float), rtol=0, atol=1e-12)
    predicted = model.predict(features)
    np.testing.assert_allclose(predicted[list(PREDICTION_COLUMNS)].astype(float),
                               old[list(PREDICTION_COLUMNS)].astype(float), rtol=0, atol=1e-9)
    base = old.copy(deep=True)
    base[list(PREDICTION_COLUMNS)] = predicted[list(PREDICTION_COLUMNS)]
    table, reference = attach_reference(base, locations, ROOT, settings["ausrap_source"])
    table = table.reindex(columns=OUTPUT_COLUMNS)
    metadata = {
        **original_metadata,
        "run_id": uuid4().hex,
        "workflow_version": settings["workflow_version"],
        "output_schema_version": OUTPUT_SCHEMA_VERSION,
        "started_at_utc": started,
        "finished_at_utc": datetime.now(UTC).isoformat(),
        "configuration": settings,
        "prediction_basis": "released_model_inference",
        "execution_mode": "saved_feature_replay",
        "feature_reuse": {
            "source": SNAPSHOT,
            "source_sha256": source_sha,
            "original_workbook_sha256": payload["source_hashes"]["predictions.xlsx"],
            "original_manifest_sha256": payload["source_hashes"]["predictions.manifest.json"],
            "original_run_id": original_metadata["run_id"],
            "original_finished_at_utc": original_metadata["finished_at_utc"],
            "external_requests_made": False,
            "note": "Context, imagery diagnostics and row warnings are retained from the original extraction. Model inference and AusRAP attachment were rerun locally.",
        },
        "ausrap": reference,
    }
    output.mkdir(parents=True, exist_ok=True)
    paths = write_prediction_outputs(PredictionResult(table, features, metadata), output)
    old[INPUT_COLUMNS].to_csv(output / "input_locations.csv", index=False)
    saved = pd.read_excel(paths["workbook"], sheet_name="Prediction_Output")
    assert tuple(saved.columns) == OUTPUT_COLUMNS and len(saved) == 10
    assert saved.site_id.tolist() == old.site_id.tolist()
    observed = ["obs_peak_cross_ped_per_h", "obs_peak_along_ped_per_h", "observed_cross_category"]
    assert saved[observed].isna().all().all()
    np.testing.assert_allclose(saved[list(PREDICTION_COLUMNS)],
                               old[list(PREDICTION_COLUMNS)].astype(float), rtol=0, atol=1e-9)
    saved_features = pd.read_excel(paths["workbook"], sheet_name="Features", dtype=object)
    # Missing cells round-trip as NaN; compare with the same representation.
    pd.testing.assert_frame_equal(saved_features, features.where(features.notna(), np.nan))
    book = load_workbook(paths["workbook"], data_only=False)
    assert [book["Metadata"].cell(row, 1).value for row in range(19, 48)] == list(OUTPUT_COLUMNS)
    assert all(cell.data_type != "e" for sheet in book for row in sheet for cell in row)
    assert sha(source) == source_sha
    verification = {
        "row_count": len(saved), "column_count": len(saved.columns),
        "model_predictions_match_saved_run_atol": 1e-9,
        "saved_features_unchanged": True, "source_snapshot_unchanged": True,
        "observations_blank": True, "accuracy_metrics_computed": False,
        "external_requests_made": False,
        "ausrap_source_rows": reference["source"]["source_rows"],
        "comparison_status_counts": reference["status_counts"],
        "prediction_status_counts": table.status.value_counts().to_dict(),
        "workbook_sha256": sha(paths["workbook"]),
    }
    (output / "verification.json").write_text(json.dumps(verification, indent=2) + "\n")
    print(json.dumps(verification, indent=2))


if __name__ == "__main__":
    main()
